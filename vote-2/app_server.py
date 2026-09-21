#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
বাংলাদেশ ভোটার তালিকা অ্যাপ -- লোকাল ওয়েব সার্ভার
==================================================
voter_extractor.py-কে সরাসরি রিইউজ করে PDF কনভার্সন, লাইভ প্রগ্রেস
(SSE) এবং এক্সট্র্যাক্ট করা ডেটা ব্রাউজ/ফিল্টার/এডিট করার API দেয়।

চালানো: python app_server.py  (তারপর ব্রাউজারে http://127.0.0.1:8765)
"""

import asyncio
import json
import os
import string
import threading
import time
import uuid
from datetime import datetime
from pathlib import Path

import openpyxl
from fastapi import FastAPI, HTTPException
from fastapi.responses import HTMLResponse, StreamingResponse
from pydantic import BaseModel

import voter_extractor as ve

APP_DIR  = Path(__file__).parent
ROOT_DIR = APP_DIR.parent  # ...\Vote -- চট্টগ্রাম-৬ ইত্যাদি ফোল্ডার এখানে থাকে

app = FastAPI(title="ভোটার তালিকা অ্যাপ")

# ──────────────────────────────────────────
# ব্যাকগ্রাউন্ড কনভার্সন জব
# ──────────────────────────────────────────
JOBS: dict = {}
JOB_LOCK = threading.Lock()
ACTIVE_JOB_ID = None


def run_conversion_job(job_id: str, pdf_path: str):
    job = JOBS[job_id]
    job["status"] = "running"
    t0 = time.time()

    def on_progress(page_idx, total_pages, cell_idx, total_cells, voter_count):
        job["progress"] = {
            "page": page_idx, "total_pages": total_pages,
            "cell": cell_idx, "total_cells": total_cells,
            "voters": voter_count, "elapsed": round(time.time() - t0, 1),
        }

    try:
        p = Path(pdf_path)
        fm = ve.meta_from_path(p, p)  # single-file মোড -- CLI-এর মতোই root=ফাইল নিজেই
        voters = ve.process_pdf(p, fm, progress_cb=on_progress)
        elapsed = time.time() - t0

        if not voters:
            job["status"] = "done"
            job["result"] = {"output_file": None, "total_voters": 0, "flagged": 0,
                              "elapsed_minutes": round(elapsed / 60, 1)}
            return

        flagged = sum(1 for v in voters if v.get("_flag"))
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = os.path.join(ve.OUTPUT_DIR, f"voter_data_{ts}.xlsx")
        stats = dict(total_voters=len(voters), total_pdfs=1, success=1, failed=0,
                     elapsed=f"{elapsed/60:.1f} মিনিট")
        ve.write_excel(voters, out_path, stats)

        job["status"] = "done"
        job["result"] = {
            "output_file": out_path,
            "total_voters": len(voters),
            "flagged": flagged,
            "elapsed_minutes": round(elapsed / 60, 1),
        }
    except Exception as e:
        job["status"] = "error"
        job["error"] = str(e)
    finally:
        global ACTIVE_JOB_ID
        with JOB_LOCK:
            if ACTIVE_JOB_ID == job_id:
                ACTIVE_JOB_ID = None


class ConvertRequest(BaseModel):
    pdf_path: str


@app.post("/api/convert")
def start_convert(payload: ConvertRequest):
    global ACTIVE_JOB_ID
    pdf_path = payload.pdf_path
    if not os.path.isfile(pdf_path) or not pdf_path.lower().endswith(".pdf"):
        raise HTTPException(400, "বৈধ PDF ফাইল নয়")

    with JOB_LOCK:
        if ACTIVE_JOB_ID and JOBS.get(ACTIVE_JOB_ID, {}).get("status") in ("starting", "running"):
            raise HTTPException(409, "ইতিমধ্যে একটি কনভার্সন চলছে -- সেটি শেষ হওয়া পর্যন্ত অপেক্ষা করুন")
        job_id = uuid.uuid4().hex[:8]
        JOBS[job_id] = {"status": "starting", "pdf_path": pdf_path,
                         "progress": None, "result": None, "error": None}
        ACTIVE_JOB_ID = job_id

    threading.Thread(target=run_conversion_job, args=(job_id, pdf_path), daemon=True).start()
    return {"job_id": job_id}


@app.get("/api/jobs/{job_id}/stream")
async def stream_job(job_id: str):
    if job_id not in JOBS:
        raise HTTPException(404, "job পাওয়া যায়নি")

    async def gen():
        last = None
        while True:
            job = JOBS.get(job_id)
            if job is None:
                break
            payload = {"status": job["status"], "progress": job["progress"]}
            if job["status"] in ("done", "error"):
                payload["result"] = job.get("result")
                payload["error"] = job.get("error")
                yield f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"
                break
            if payload != last:
                yield f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"
                last = payload
            await asyncio.sleep(0.4)

    return StreamingResponse(gen(), media_type="text/event-stream")


# ──────────────────────────────────────────
# ফোল্ডার ব্রাউজার (PDF বাছাই করার জন্য)
# ──────────────────────────────────────────
@app.get("/api/browse")
def browse(path: str = ""):
    if not path:
        roots = [str(ROOT_DIR), os.path.normpath(ve.OUTPUT_DIR)]
        try:
            roots.extend(os.listdrives())
        except AttributeError:
            for letter in string.ascii_uppercase:
                d = f"{letter}:\\"
                if os.path.exists(d):
                    roots.append(d)
        seen, entries = set(), []
        for r in roots:
            if r not in seen:
                seen.add(r)
                entries.append({"name": r, "path": r, "type": "dir"})
        return {"path": "", "parent": None, "entries": entries}

    p = Path(path)
    if not p.exists() or not p.is_dir():
        raise HTTPException(404, "ফোল্ডার পাওয়া যায়নি")

    entries = []
    try:
        for child in sorted(p.iterdir(), key=lambda c: (not c.is_dir(), c.name.lower())):
            if child.is_dir():
                entries.append({"name": child.name, "path": str(child), "type": "dir"})
            elif child.suffix.lower() == ".pdf":
                entries.append({"name": child.name, "path": str(child), "type": "pdf",
                                 "size_mb": round(child.stat().st_size / 1024 / 1024, 2)})
    except PermissionError:
        pass

    parent = str(p.parent) if p.parent != p else None
    return {"path": str(p), "parent": parent, "entries": entries}


# ──────────────────────────────────────────
# আগে কনভার্ট করা Excel ফাইলের তালিকা
# ──────────────────────────────────────────
@app.get("/api/outputs")
def list_outputs():
    out_dir = Path(ve.OUTPUT_DIR)
    if not out_dir.exists():
        return []
    files = sorted(out_dir.glob("voter_data_*.xlsx"), key=lambda f: f.stat().st_mtime, reverse=True)
    return [{"path": str(f), "name": f.name, "mtime": f.stat().st_mtime} for f in files]


# ──────────────────────────────────────────
# ভোটার ডেটা লোড
# ──────────────────────────────────────────
@app.get("/api/voters")
def get_voters(file: str):
    if not os.path.isfile(file):
        raise HTTPException(404, "ফাইল পাওয়া যায়নি")
    wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    if "ভোটার তালিকা" not in wb.sheetnames:
        raise HTTPException(400, "এটি একটি ভোটার তালিকা Excel ফাইল নয়")
    ws = wb["ভোটার তালিকা"]

    rows = []
    for ri, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if row is None or all(c is None for c in row):
            continue
        rec = {"_row": ri}
        for ci, key in enumerate(ve.COL_KEYS):
            val = row[ci] if ci < len(row) and row[ci] is not None else ""
            rec[key] = val
        rows.append(rec)
    wb.close()

    columns = [{"key": key, "label": label} for key, label, _ in ve.COLS]
    return {"columns": columns, "rows": rows}


# ──────────────────────────────────────────
# এডিট সেভ
# ──────────────────────────────────────────
class EditItem(BaseModel):
    row: int
    key: str
    value: str


class SaveRequest(BaseModel):
    file: str
    edits: list[EditItem]


@app.post("/api/voters/save")
def save_voters(payload: SaveRequest):
    if not os.path.isfile(payload.file):
        raise HTTPException(404, "ফাইল পাওয়া যায়নি")
    if not payload.edits:
        return {"status": "ok", "updated_rows": []}

    wb = openpyxl.load_workbook(payload.file)
    ws = wb["ভোটার তালিকা"]
    col_index = {key: i + 1 for i, key in enumerate(ve.COL_KEYS)}

    touched_rows = set()
    for edit in payload.edits:
        ci = col_index.get(edit.key)
        if not ci or edit.key == "_flag":
            continue
        ws.cell(row=edit.row, column=ci, value=edit.value)
        touched_rows.add(edit.row)

    for ri in touched_rows:
        v = {key: (ws.cell(row=ri, column=ci).value or "") for key, ci in col_index.items() if key != "_flag"}
        flag_str = "; ".join(ve.flag_record(v))
        ws.cell(row=ri, column=col_index["_flag"], value=flag_str)
        is_flagged = bool(flag_str)
        for key, ci in col_index.items():
            ve.style_data_cell(ws.cell(row=ri, column=ci), ci, ri, is_flagged)

    wb.save(payload.file)
    return {"status": "ok", "updated_rows": sorted(touched_rows)}


# ──────────────────────────────────────────
# ফ্রন্টএন্ড
# ──────────────────────────────────────────
@app.get("/", response_class=HTMLResponse)
def index():
    return (APP_DIR / "voter_app.html").read_text(encoding="utf-8")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="127.0.0.1", port=8765)
