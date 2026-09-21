"""PDF কনভার্সনের পর একটা ব্যাকআপ Excel ফাইল সেভ করার জন্য (DB-তে upsert করার পাশাপাশি)।"""
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .voter_rules import COLS

HEADER_FILL = PatternFill("solid", start_color="162B4D")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=10)
FLAG_FILL = PatternFill("solid", start_color="FDE2E2")
FLAG_FONT = Font(name="Arial", size=9, color="B00020")
DATA_FONT = Font(name="Arial", size=9)
THIN = Side(style="thin", color="DDDDDD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_excel_backup(voters: list[dict], out_path: str):
    """voters -- বাংলা কী সহ dict-এর লিস্ট (ocr_engine.process_pdf()-এর আউটপুট)"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ভোটার তালিকা"

    for ci, (_, label, _field) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=ci, value=label)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        ws.column_dimensions[get_column_letter(ci)].width = 20

    for ri, v in enumerate(voters, 2):
        is_flagged = bool(v.get('_flag'))
        for ci, (bkey, _label, _field) in enumerate(COLS, 1):
            c = ws.cell(row=ri, column=ci, value=v.get(bkey, ''))
            c.border = BORDER
            c.font = FLAG_FONT if is_flagged else DATA_FONT
            if is_flagged:
                c.fill = FLAG_FILL

    ws.freeze_panes = "A2"
    wb.save(out_path)
