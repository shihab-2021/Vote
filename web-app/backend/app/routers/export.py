import csv
import io

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, Request, status
from fastapi.responses import StreamingResponse
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from ..audit import log_activity
from ..auth import require_permission
from ..db import get_db
from ..models import User, Voter
from ..voter_rules import COLS
from .voters import build_voter_query

router = APIRouter(prefix="/api/export", tags=["export"])

EXPORT_COLS = [(field, label) for _, label, field in COLS if field] + [("flag_reasons", "যাচাই প্রয়োজন")]
MAX_EXPORT_ROWS = 5000  # একটা কুয়েরিতে অনির্দিষ্ট আকারের ফাইল তৈরি হয়ে যাওয়া আটকাতে


def _filtered_voters(
    db: Session, user: User, search: str, ward: str, upazila: str, flagged: bool | None, filters: str = "",
):
    q = build_voter_query(db, user, search=search, ward=ward, upazila=upazila, flagged=flagged, filters=filters)
    q = q.order_by(Voter.name)
    return db.scalars(q).all()


def _row_value(voter: Voter, field: str):
    if field == "flag_reasons":
        return "; ".join(voter.flag_reasons or [])
    return getattr(voter, field, "") or ""


@router.get("")
def export_voters(
    request: Request,
    format: str = "xlsx",
    search: str = "", ward: str = "", upazila: str = "",
    flagged: bool | None = None,
    filters: str = "",
    db: Session = Depends(get_db),
    user: User = Depends(require_permission("export_voter")),
):
    voters = _filtered_voters(db, user, search, ward, upazila, flagged, filters)
    if len(voters) > MAX_EXPORT_ROWS:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            f"একবারে সর্বোচ্চ {MAX_EXPORT_ROWS} জন এক্সপোর্ট করা যাবে -- ফিল্টার আরও নির্দিষ্ট করুন "
            f"({len(voters)} জন মিলেছে)",
        )

    # লগ-অ্যাক্টিভিটির db.commit() সবসময় ফাইল তৈরি হওয়ার *পরে* কল হয় -- SQLAlchemy সেশনের
    # ডিফল্ট expire_on_commit=True আচরণের কারণে commit-এর পর voters লিস্টের অবজেক্টগুলো "expired"
    # হয়ে যায়, ফলে commit আগে করলে নিচের লুপে প্রতিটা attribute access-এ আলাদা করে DB থেকে
    # রিলোড হতো (হাজার হাজার voter-এর জন্য কার্যত অসীম ধীরগতি -- আসলে ঘটেছিল, তাই এই কমেন্ট)।
    if format == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow([label for _, label in EXPORT_COLS])
        for v in voters:
            writer.writerow([_row_value(v, field) for field, _ in EXPORT_COLS])
        content = buf.getvalue()

        log_activity(db, user, "export_voters", detail={"format": format, "count": len(voters)}, request=request)
        db.commit()
        return StreamingResponse(
            iter([content]), media_type="text/csv",
            headers={"Content-Disposition": "attachment; filename=voters_export.csv"},
        )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ভোটার তালিকা"

    header_fill = PatternFill("solid", start_color="162B4D")
    header_font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    thin = Side(style="thin", color="DDDDDD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    flag_fill = PatternFill("solid", start_color="FDE2E2")
    flag_font = Font(name="Arial", size=9, color="B00020")
    data_font = Font(name="Arial", size=9)

    for ci, (_, label) in enumerate(EXPORT_COLS, 1):
        c = ws.cell(row=1, column=ci, value=label)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
        ws.column_dimensions[get_column_letter(ci)].width = 20

    for ri, voter in enumerate(voters, 2):
        is_flagged = bool(voter.flag_reasons)
        for ci, (field, _) in enumerate(EXPORT_COLS, 1):
            c = ws.cell(row=ri, column=ci, value=_row_value(voter, field))
            c.border = border
            c.font = flag_font if is_flagged else data_font
            if is_flagged:
                c.fill = flag_fill

    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    log_activity(db, user, "export_voters", detail={"format": format, "count": len(voters)}, request=request)
    db.commit()
    return StreamingResponse(
        out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=voters_export.xlsx"},
    )
