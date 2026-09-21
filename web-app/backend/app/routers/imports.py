import io

import openpyxl
from fastapi import APIRouter, Depends, HTTPException, UploadFile, status
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..auth import get_current_user, require_admin
from ..db import get_db
from ..import_logic import VoterUpserter, bengali_row_to_record
from ..models import ImportBatch, User
from ..schemas import ImportBatchOut, ImportResult
from ..voter_rules import BENGALI_KEYS

router = APIRouter(prefix="/api/import", tags=["import"])


@router.get("/batches", response_model=list[ImportBatchOut])
def list_batches(db: Session = Depends(get_db), _user: User = Depends(get_current_user)):
    return db.scalars(select(ImportBatch).order_by(ImportBatch.imported_at.desc())).all()


@router.post("", response_model=ImportResult)
def import_excel(
    file: UploadFile,
    db: Session = Depends(get_db),
    user: User = Depends(require_admin),
):
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "শুধু .xlsx/.xls ফাইল সমর্থিত")

    content = file.file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Excel ফাইল পড়া যায়নি -- ফাইলটি করাপ্ট হতে পারে")

    if "ভোটার তালিকা" not in wb.sheetnames:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "এটি voter_extractor.py-এর তৈরি Excel ফাইল নয়")
    ws = wb["ভোটার তালিকা"]

    batch = ImportBatch(filename=file.filename, imported_by=user.id, status="processing")
    db.add(batch)
    db.flush()

    errors = total_rows = 0
    upserter = VoterUpserter(db)

    try:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row is None or all(c is None for c in row):
                continue
            total_rows += 1
            try:
                v = {bkey: (str(row[ci]).strip() if ci < len(row) and row[ci] is not None else "")
                     for ci, bkey in enumerate(BENGALI_KEYS)}
                record = bengali_row_to_record(v, batch.id, user.id)
                if record is None:
                    errors += 1
                    continue
                upserter.add(record)
            except Exception:
                errors += 1
                continue

        upserter.flush()

        batch.row_count = total_rows
        batch.inserted_count = upserter.inserted
        batch.updated_count = upserter.updated
        batch.error_count = errors
        batch.status = "done"
        db.commit()
    except Exception as e:
        db.rollback()
        batch.status = "failed"
        db.add(batch)
        db.commit()
        raise HTTPException(status.HTTP_500_INTERNAL_SERVER_ERROR, f"ইমপোর্ট ব্যর্থ: {e}")

    return ImportResult(
        import_batch_id=batch.id, inserted=upserter.inserted, updated=upserter.updated,
        errors=errors, total_rows=total_rows,
    )
